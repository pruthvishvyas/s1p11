import os, sys, json, datetime
sys.path.append(os.path.abspath(os.path.join(os.path.dirname(__file__), '../../')))
from config.config import CONFIG, save_json, safe_div, compute_thresholds, rel
import numpy as np
import pandas as pd

WS_COLS = {'Week_Start': 'Week Start', 'Promo_Event': 'Promo', 'Total_Marketing_Spend': 'Marketing Spend', 'Orders': 'Orders',
           'Net_Revenue': 'Net Revenue', 'Net_Profit': 'Net Profit', 'Net_Margin_pct': 'Net Margin %', 'ROAS': 'ROAS', 'CAC': 'CAC',
           'LTV_to_CAC': 'LTV to CAC', 'segment': 'Week regime', 'is_anomaly': 'Unusual week', 'loss_risk_flag': 'Loss risk flag'}


# ------------------------------------------------------------------ Excel
def _style(ws):
    from openpyxl.styles import PatternFill, Font, Alignment
    fill, font = PatternFill('solid', fgColor='1565C0'), Font(bold=True, color='FFFFFF')
    for c in ws[1]:
        c.fill, c.font, c.alignment = fill, font, Alignment(wrap_text=True, vertical='center')
    ws.freeze_panes = 'A2'
    for col in ws.columns:
        width = max(len(str(c.value)) if c.value is not None else 0 for c in list(col)[:60])
        ws.column_dimensions[col[0].column_letter].width = min(max(10, width + 2), 60)


def _excel(config, df, ctx):
    path = config['EXCEL_PATH']
    os.makedirs(os.path.dirname(path), exist_ok=True)
    d = df.copy()
    d[config['DATE_COL']] = pd.to_datetime(d[config['DATE_COL']]).dt.date
    ws = d[[c for c in WS_COLS if c in d.columns]].rename(columns=WS_COLS).iloc[::-1].round(2)
    ch = ctx.get('seg_priority')
    if ch is None or len(ch) == 0:
        ch = ctx.get('agg_channel', pd.DataFrame())
    ch = ch.rename(columns={c: c.replace('_', ' ').capitalize() for c in ch.columns}).round(3)
    pa = ctx.get('agg_promo', pd.DataFrame())
    pa = pa.rename(columns={c: c.replace('_', ' ').capitalize() for c in pa.columns}).round(2)
    cash = d[[config['DATE_COL'], 'Cash_Balance', 'AR_Outstanding', 'DSO_Days', 'cash_runway_weeks']].rename(
        columns={config['DATE_COL']: 'Week Start', 'Cash_Balance': 'Cash Balance', 'AR_Outstanding': 'AR Outstanding',
                 'DSO_Days': 'DSO Days', 'cash_runway_weeks': 'Cash runway (weeks)'}).iloc[::-1].round(2)
    fc = ctx.get('forecast_df', pd.DataFrame()).copy()
    if len(fc) and config['DATE_COL'] in fc:
        fc[config['DATE_COL']] = pd.to_datetime(fc[config['DATE_COL']]).dt.date
    fc = fc.rename(columns={c: c.replace('_', ' ') for c in fc.columns}).round(2)
    rows = []
    for r in ctx.get('recommendations', []):
        rows.append({'Priority': r['priority'], 'Area': r['scope'].title(), 'Subject': r['subject'], 'Action': r['action'], 'Evidence': r['evidence']})
    for i in ctx.get('insights', []):
        rows.append({'Priority': i['severity'], 'Area': i['category'], 'Subject': i['finding'], 'Action': i['action'], 'Evidence': i['evidence']})
    order = {'HIGH': 0, 'MEDIUM': 1, 'LOW': 2}
    actions = pd.DataFrame(rows)
    if len(actions):
        actions = actions.sort_values('Priority', key=lambda s: s.map(order)).reset_index(drop=True)
    with pd.ExcelWriter(path, engine='openpyxl') as xw:
        for name, frame in [('Weekly Summary', ws), ('Channel Performance', ch), ('Promo Analysis', pa),
                            ('Cash & Collections', cash), ('13-Week Outlook', fc), ('Action List', actions)]:
            (frame if len(frame) else pd.DataFrame({'Note': ['No data for this sheet in this run']})).to_excel(xw, sheet_name=name, index=False)
            _style(xw.sheets[name])
    return path


def _html(config, df, ctx):
    try:
        import plotly.graph_objects as go
    except Exception:
        return None
    try:
        fig = go.Figure()
        fig.add_trace(go.Scatter(x=df[config['DATE_COL']], y=df['Net_Revenue'], name='Net Revenue'))
        fig.add_trace(go.Scatter(x=df[config['DATE_COL']], y=df['Net_Profit'], name='Net Profit'))
        fc = ctx.get('forecast_df')
        if fc is not None and len(fc):
            fig.add_trace(go.Scatter(x=fc[config['DATE_COL']], y=fc['Net_Revenue_forecast'], name='Net Revenue forecast', line=dict(dash='dash')))
            fig.add_trace(go.Scatter(x=fc[config['DATE_COL']], y=fc['Net_Profit_forecast'], name='Net Profit forecast', line=dict(dash='dash')))
        fig.update_layout(title='Weekly Growth & Profit with 13-week outlook', template='plotly_white')
        fig.write_html(config['HTML_PATH'], include_plotlyjs='cdn')
        return config['HTML_PATH']
    except Exception as e:
        print(f"  HTML dashboard skipped: {e}")
        return None


# ------------------------------------------------------------------ frontend export
def _kpis(df, config):
    H = config['BASELINE_WINDOW_WEEKS']
    a, b = df.tail(H), df.iloc[-2 * H:-H]

    def entry(key, label, v, pv, unit, period):
        ch = ((v - pv) / abs(pv) * 100) if (pv is not None and pv != 0) else None
        return {'key': key, 'label': label, 'value': v, 'prior_value': pv, 'change_pct': ch, 'unit': unit, 'period': period}
    pw = f'last {H} weeks'
    rev, prev = float(a['Net_Revenue'].sum()), float(b['Net_Revenue'].sum()) if len(b) else None
    prof, prof_p = float(a['Net_Profit'].sum()), float(b['Net_Profit'].sum()) if len(b) else None
    marg = float(safe_div(a['Net_Profit'].sum(), a['Net_Revenue'].sum()) * 100)
    marg_p = float(safe_div(b['Net_Profit'].sum(), b['Net_Revenue'].sum()) * 100) if len(b) else None
    roas = float(safe_div(a['Net_Revenue'].sum(), a[config['TOTAL_SPEND']].sum()))
    roas_p = float(safe_div(b['Net_Revenue'].sum(), b[config['TOTAL_SPEND']].sum())) if len(b) else None
    cash_p = float(df['Cash_Balance'].iloc[-H - 1]) if len(df) > H else None
    run_p = float(df['cash_runway_weeks'].iloc[-H - 1]) if len(df) > H else None
    k = [entry('net_revenue', 'Net Revenue', rev, prev, 'currency', pw),
         entry('net_profit', 'Net Profit', prof, prof_p, 'currency', pw),
         entry('net_margin_pct', 'Net Margin %', marg, marg_p, 'percent', pw),
         entry('roas', 'ROAS', roas, roas_p, 'ratio', pw),
         entry('cac', 'CAC', float(a['CAC'].mean()), float(b['CAC'].mean()) if len(b) else None, 'currency', pw),
         entry('ltv_to_cac', 'LTV to CAC', float(a['LTV_to_CAC'].mean()), float(b['LTV_to_CAC'].mean()) if len(b) else None, 'ratio', pw),
         entry('cash_balance', 'Cash Balance', float(df['Cash_Balance'].iloc[-1]), cash_p, 'currency', 'latest week'),
         entry('cash_runway_weeks', 'Cash runway (weeks)', float(df['cash_runway_weeks'].iloc[-1]), run_p, 'weeks', 'latest week'),
         entry('loss_week_count', 'Loss weeks', int(df[config['TARGET_FLAG']].sum()), None, 'count', 'all history'),
         entry('anomaly_count', 'Unusual weeks', int(df['is_anomaly'].sum()) if 'is_anomaly' in df else 0, None, 'count', 'all history')]
    k[8]['last_52_weeks'] = int(df[config['TARGET_FLAG']].tail(52).sum())
    return k


def _push_bigquery(config, df):
    bq = config['BIGQUERY']
    if not bq.get('project_id') or not bq.get('dataset'):
        return False, 'BigQuery not configured; CSV fallback used'
    try:
        from google.cloud import bigquery
        from google.oauth2 import service_account
        cred = bq['credentials_file']
        cred = cred if os.path.isabs(cred) else os.path.join(config['ROOT'], cred)
        creds = service_account.Credentials.from_service_account_file(cred)
        client = bigquery.Client(project=bq['project_id'], credentials=creds)
        table_id = f"{bq['project_id']}.{bq['dataset']}.{bq['table']}"
        job = client.load_table_from_dataframe(df, table_id, job_config=bigquery.LoadJobConfig(write_disposition='WRITE_TRUNCATE'))
        job.result()
        return True, table_id
    except Exception as e:
        return False, f'BigQuery failed ({e}); CSV fallback used'


def _read_json(path, default):
    try:
        with open(path, 'r', encoding='utf-8') as f:
            return json.load(f)
    except Exception:
        return default


def run_export(config=CONFIG, df_feat=None, insights=None, **ctx):
    """Phase 10-11: Excel report, BigQuery-first / CSV fallback, output/ JSON for the frontend. Outputs: excel_path, html_path, output_dir, source_type, kpis."""
    if df_feat is None:
        df_feat = pd.read_csv(config['FEAT_CSV'], parse_dates=[config['DATE_COL']])
    if insights is None:
        insights = _read_json(config['INSIGHTS_JSON'], [])
    df = df_feat.copy().sort_values(config['DATE_COL']).reset_index(drop=True)
    df[config['DATE_COL']] = pd.to_datetime(df[config['DATE_COL']])
    ctx['insights'] = insights
    thr = ctx.get('thresholds') or compute_thresholds(config, df)
    out = config['OUTPUT_DIR']
    j = lambda *p: os.path.join(out, *p)

    excel_path = _excel(config, df, ctx)
    html_path = _html(config, df, ctx)
    ok, msg = _push_bigquery(config, df)
    source_type = 'bigquery' if ok else 'csv'
    print(f"  Data source: {source_type} ({msg})")

    kpis = _kpis(df, config)
    save_json(j('understand', 'kpis.json'), {'kpis': kpis, 'thresholds': thr})
    seg_report = ctx.get('seg_report', {})
    prof = ctx.get('cluster_profile')
    seg_cols = [c for c in [config['ID_COL'], config['DATE_COL'], 'segment', 'segment_tier'] if c in df.columns]
    save_json(j('understand', 'segments.json'), {'report': seg_report, 'profile': prof if prof is not None else [],
                                                 'assignments': df[seg_cols] if seg_cols else []})
    an = df[df['is_anomaly'] == 1] if 'is_anomaly' in df else pd.DataFrame()
    an_cols = [c for c in [config['ID_COL'], config['DATE_COL'], 'anomaly_score', 'anomaly_reason', 'Net_Profit', 'Net_Revenue'] if c in an.columns]
    save_json(j('understand', 'anomalies.json'), {'count': int(len(an)), 'contamination': config['ANOMALY_CONT'],
                                                  'items': an[an_cols].sort_values('anomaly_score', ascending=False) if len(an) else []})

    goals = ctx.get('goals') or _read_json(config['GOALS_JSON'], config['GOAL_DEFAULTS'])
    save_json(j('decide', 'insights.json'), {'insights': insights})
    save_json(j('decide', 'recommendations.json'), {'recommendations': ctx.get('recommendations', []),
                                                    'channel_priority': ctx.get('seg_priority', []),
                                                    'weekly_flags_csv': rel(config['BL_CSV'], config)})
    for name, obj in (ctx.get('simulations') or {}).items():
        save_json(j('decide', 'simulations', f'{name}.json'), obj)
    save_json(j('decide', 'goals.json'), goals)

    fc = ctx.get('forecast_df')
    save_json(j('monitor', 'forecast.json'), {'report': ctx.get('forecast_report', {}), 'forecast': fc if fc is not None else []})
    save_json(j('monitor', 'goal_progress.json'), ctx.get('goal_progress', {}))
    save_json(j('monitor', 'alerts.json'), {'alerts': ctx.get('alerts', []), 'count': len(ctx.get('alerts', []))})

    clf = ctx.get('clf_report', {})
    drv = ctx.get('driver_report', {})
    fcr = ctx.get('forecast_report', {})
    val = ctx.get('validation_dict', _read_json(config['VALIDATION_JSON'], {}))
    summary = {'dataset': config['DATASET_NAME'], 'client_id': config['CLIENT_ID'], 'pipeline_version': config['PIPELINE_VERSION'],
               'generated_at': datetime.datetime.now().isoformat(timespec='seconds'), 'rows': int(len(df)),
               'date_range': [df[config['DATE_COL']].min().strftime('%Y-%m-%d'), df[config['DATE_COL']].max().strftime('%Y-%m-%d')],
               'synthetic_data': bool(val.get('synthetic_data', False)), 'roas_basis': val.get('roas_basis'),
               'promo_weeks': int(df['promo_flag'].sum()), 'loss_weeks': int(df[config['TARGET_FLAG']].sum()),
               'anomaly_weeks': int(len(an)), 'thresholds': thr, 'source_type': source_type,
               'models': {'segmentation': seg_report, 'loss_week_classifier': clf, 'driver_model': drv,
                          'forecast': {t: {k: v for k, v in d.items() if k in ('method', 'model_accepted', 'best_baseline', 'shipped_mape_pct', 'mape_target_met', 'accuracy_claim', 'next_13w_total', 'last_13w_actual_total', 'model_metrics', 'seasonal_naive_metrics', 'moving_avg_4w_metrics', 'note')}
                                       for t, d in (fcr.get('targets') or {}).items()},
                          'forecast_direction': fcr.get('direction'), 'next_week_loss_risk': fcr.get('next_week_loss_risk'),
                          'forecast_assumptions': fcr.get('assumptions')},
               'files': {'excel': rel(excel_path, config), 'feature_store': rel(config['FEAT_CSV'], config)}}
    save_json(j('meta', 'summary.json'), summary)
    save_json(j('meta', 'terminology.json'), config['TERMINOLOGY'])
    log_path = j('meta', 'run_log.json')
    log = _read_json(log_path, [])
    if not isinstance(log, list):
        log = []
    log.append({'timestamp': summary['generated_at'], 'pipeline_version': config['PIPELINE_VERSION'], 'rows': summary['rows'],
                'source_type': source_type, 'synthetic_data': summary['synthetic_data'],
                'silhouette': seg_report.get('silhouette'), 'loss_week_method': clf.get('method'),
                'forecast_methods': {t: d.get('method') for t, d in (fcr.get('targets') or {}).items()},
                'alerts': len(ctx.get('alerts', [])), 'insights': len(insights)})
    save_json(log_path, log)

    contract = {
        'project': {'name': 'Weekly Growth & Profit', 'domain': 'd2c_ecommerce', 'client_id': config['CLIENT_ID'],
                    'description': 'Weekly channel, promo and cash analytics'},
        'terminology': config['TERMINOLOGY'],
        'outcomes': {
            'understand': {'kpis': 'output/understand/kpis.json', 'charts': 'output/understand/charts/',
                           'segments': 'output/understand/segments.json', 'anomalies': 'output/understand/anomalies.json'},
            'decide': {'insights': 'output/decide/insights.json', 'recommendations': 'output/decide/recommendations.json',
                       'simulations': 'output/decide/simulations/', 'goals': 'output/decide/goals.json'},
            'monitor': {'forecast': 'output/monitor/forecast.json', 'goal_progress': 'output/monitor/goal_progress.json',
                        'alerts': 'output/monitor/alerts.json'}},
        'model': {'type': 'GradientBoostingRegressor (or seasonal-naive if baseline wins)', 'pkl_file': 'models/forecast_model.pkl',
                  'input_fields': ['Total_Marketing_Spend', 'Spend_Search', 'Spend_Social', 'Spend_Email', 'Spend_Display',
                                   'promo_flag', 'discount_rate_pct', 'week_of_year'],
                  'output_label': 'Net_Revenue', 'horizon': '13 weeks'},
        'source': {'type': source_type, 'file': 'data/processed/d2c_weekly_growth_feature_store.csv'},
        'frontend': {'framework': 'vanilla HTML/JS', 'deploy_target': 'cloudflare_pages'},
        'saas_ready': True}
    save_json(j('frontend_contract.json'), contract)
    save_json(os.path.join(config['ROOT'], 'frontend_contract.json'), contract)
    print(f"  Excel: {excel_path}")
    print(f"  Frontend files written under {out}")
    return {'excel_path': excel_path, 'html_path': html_path, 'output_dir': out, 'source_type': source_type, 'kpis': kpis}

